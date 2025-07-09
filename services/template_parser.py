import os
import re
import json
from docx import Document
import google.generativeai as genai
from config import GEMINI_API_KEY

# Configure Gemini AI
genai.configure(api_key=GEMINI_API_KEY)
MODEL = genai.GenerativeModel("gemini-1.5-flash")

def extract_text_from_docx(docx_path):
    """Extract text content from a DOCX file"""
    doc = Document(docx_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)

def build_prompt(template_text):
    """Build the prompt for Gemini AI to extract schema"""
    return f"""
You are a document template parser.

Your task is to convert any given document template into a JSON schema with:
- Field metadata (ID, label, type)
- A templateString using {{placeholders}} where data goes

👀 Detect dynamic fields such as:
- Underlines (_________), placeholders ([Date], [Company Name])
- Table headers with repeating data (e.g., Item Name, Quantity, Price)

If you see a table-like structure, define it as:
"type": "array of objects" and include its "itemSchema"

🎯 OUTPUT FORMAT:
{{
  "name": "<Template Name>",
  "fields": [...],
  "templateString": "..."
}}

🎓 Field Types:
- string
- date
- number
- array of objects

📌 Only return JSON. Do NOT wrap in ```json.

---

Document Template:

{template_text}

---
"""

def clean_and_parse_gemini_json(text):
    """Clean and parse JSON response from Gemini"""
    try:
        # Remove any markdown formatting
        cleaned = re.sub(r"^```json\s*|\s*```$", "", text.strip(), flags=re.IGNORECASE)
        return json.loads(cleaned)
    except Exception as e:
        raise ValueError(f"Failed to parse Gemini output: {e}")

def extract_schema_from_docx(docx_path):
    """Extract schema from DOCX template using Gemini AI"""
    try:
        template_text = extract_text_from_docx(docx_path)
        prompt = build_prompt(template_text)

        response = MODEL.generate_content(prompt)

        parsed = clean_and_parse_gemini_json(response.text)
        print("✅ Extracted schema:")
        print(json.dumps(parsed, indent=2))
        return parsed
    except Exception as e:
        print("❌ Error:", e)
        print("📄 Raw Gemini output:\n", getattr(response, 'text', 'No response'))
        return None